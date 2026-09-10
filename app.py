import csv
import io
import re
import smtplib
import socket
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, time as dt_time

from flask import Flask, jsonify, request, send_from_directory, send_file
from openpyxl import load_workbook, Workbook
import dns.resolver

app = Flask(__name__, static_folder="static")

jobs = {}
jobs_lock = threading.Lock()
email_re = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# Keep this list conservative. Add company-approved disposable domains as needed.
DISPOSABLE_DOMAINS = {
    "mailinator.com", "guerrillamail.com", "10minutemail.com",
    "temp-mail.org", "tempmail.com", "yopmail.com", "getnada.com",
    "trashmail.com", "sharklasers.com"
}


def normalize_email(value):
    if value is None:
        return ""
    return str(value).strip().lower()


def mx_lookup(domain):
    try:
        answers = dns.resolver.resolve(domain, "MX", lifetime=5)
        records = sorted(
            [(int(r.preference), str(r.exchange).rstrip(".")) for r in answers],
            key=lambda x: x[0]
        )
        return records
    except Exception:
        # Some domains accept mail directly on A/AAAA when no MX exists.
        try:
            socket.gethostbyname(domain)
            return [(0, domain)]
        except Exception:
            return []


def smtp_check(email, mx_host, timeout=8):
    """Best-effort SMTP RCPT check. Providers may intentionally return inconclusive results."""
    try:
        with smtplib.SMTP(timeout=timeout) as smtp:
            smtp.connect(mx_host, 25)
            smtp.ehlo_or_helo_if_needed()
            try:
                if smtp.has_extn("starttls"):
                    smtp.starttls()
                    smtp.ehlo_or_helo_if_needed()
            except Exception:
                # Continue without TLS if the server rejects it.
                pass
            smtp.mail("verify@example.com")
            code, msg = smtp.rcpt(email)
            if 200 <= code < 300:
                return "VALID", "SMTP accepted recipient"
            if 500 <= code < 600:
                return "INVALID", f"SMTP rejected recipient ({code})"
            return "UNKNOWN", f"SMTP inconclusive ({code})"
    except (socket.timeout, TimeoutError):
        return "UNKNOWN", "SMTP timeout"
    except (ConnectionRefusedError, OSError) as e:
        return "UNKNOWN", f"SMTP unavailable: {type(e).__name__}"
    except Exception as e:
        return "UNKNOWN", f"SMTP unavailable: {type(e).__name__}"


def verify_email(email, do_smtp=True):
    email = normalize_email(email)
    if not email:
        return "INVALID", "Empty email"

    if not email_re.match(email):
        return "INVALID", "Invalid email format"

    local, domain = email.rsplit("@", 1)

    if len(email) > 254 or len(local) > 64:
        return "INVALID", "Email length is invalid"

    if domain in DISPOSABLE_DOMAINS:
        return "RISKY", "Disposable/temporary email domain"

    mx = mx_lookup(domain)
    if not mx:
        return "INVALID", "Domain has no reachable MX/A record"

    if not do_smtp:
        return "VALID", f"Mail domain has MX: {mx[0][1]}"

    # Try up to two MX servers.
    last_reason = "Domain accepts DNS mail but SMTP mailbox check was inconclusive"
    for _, host in mx[:2]:
        status, reason = smtp_check(email, host)
        if status in ("VALID", "INVALID"):
            return status, reason
        last_reason = reason

    return "UNKNOWN", last_reason


def process_job(job_id, emails, do_smtp):
    total = len(emails)
    ordered = [None] * total

    with jobs_lock:
        jobs[job_id]["status"] = "running"

    def work(item):
        idx, email = item
        status, reason = verify_email(email, do_smtp)
        return idx, email, status, reason

    # Conservative concurrency. Increase only after confirming your network/provider permits it.
    with ThreadPoolExecutor(max_workers=8) as pool:
        futures = [pool.submit(work, item) for item in enumerate(emails)]
        for future in as_completed(futures):
            idx, email, status, reason = future.result()
            row = {
                "email": email,
                "status": status,
                "reason": reason
            }
            ordered[idx] = row
            with jobs_lock:
                job = jobs[job_id]
                job["done"] += 1
                job["counts"][status] = job["counts"].get(status, 0) + 1
                job["results"].append(row)
                job["ordered"] = ordered

    with jobs_lock:
        jobs[job_id]["status"] = "completed"
        jobs[job_id]["ordered"] = ordered


def excel_cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (date, dt_time, int, float, bool, str)):
        return value
    return str(value)


def copy_row(row):
    if row is None:
        return []
    return [excel_cell(c) for c in row]


def csv_cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, dt_time)):
        return value.isoformat()
    return str(value)


def find_email_column(headers):
    for i, h in enumerate(headers):
        name = str(h).strip().lower() if h is not None else ""
        if name in ("email", "e-mail", "emails", "mail", "email address", "email_address"):
            return i
    # Fallback: any header that contains "email"
    for i, h in enumerate(headers):
        name = str(h).strip().lower() if h is not None else ""
        if "email" in name:
            return i
    return None


@app.get("/")
def index():
    return send_from_directory("static", "index.html")


@app.post("/api/verify")
def start_verify():
    uploaded = request.files.get("file")
    if not uploaded:
        return jsonify({"error": "Please upload an XLSX or CSV file"}), 400

    do_smtp = request.form.get("smtp", "true").lower() == "true"
    filename = (uploaded.filename or "").lower()

    try:
        data = uploaded.read()
        if filename.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
            if not rows:
                return jsonify({"error": "Excel file is empty"}), 400
            headers = [str(x).strip() if x is not None else "" for x in rows[0]]
            email_idx = find_email_column(headers)
            if email_idx is None:
                return jsonify({"error": "Email column not found. Add a column named Email."}), 400
            source_rows = [copy_row(r) for r in rows[1:]]
            emails = [normalize_email(r[email_idx] if email_idx < len(r) else "") for r in source_rows]
            file_ext = "xlsx"
        elif filename.endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
            if not rows:
                return jsonify({"error": "CSV file is empty"}), 400
            headers = [str(x).strip() if x is not None else "" for x in rows[0]]
            email_idx = find_email_column(headers)
            if email_idx is None:
                return jsonify({"error": "Email column not found. Add a column named Email."}), 400
            source_rows = [copy_row(r) for r in rows[1:]]
            emails = [normalize_email(r[email_idx] if email_idx < len(r) else "") for r in source_rows]
            file_ext = "csv"
        else:
            return jsonify({"error": "Only XLSX and CSV files are supported"}), 400
    except Exception as e:
        return jsonify({"error": f"Could not read file: {e}"}), 400

    if not emails:
        return jsonify({"error": "No email rows found in the file"}), 400

    job_id = uuid.uuid4().hex
    with jobs_lock:
        jobs[job_id] = {
            "status": "queued",
            "total": len(emails),
            "done": 0,
            "results": [],
            "ordered": [],
            "counts": {"VALID": 0, "INVALID": 0, "RISKY": 0, "UNKNOWN": 0},
            "headers": headers,
            "source_rows": source_rows,
            "file_ext": file_ext,
        }
    threading.Thread(target=process_job, args=(job_id, emails, do_smtp), daemon=True).start()

    return jsonify({"job_id": job_id, "total": len(emails)})


@app.get("/api/status/<job_id>")
def status(job_id):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            return jsonify({"error": "Job not found"}), 404
        # Only return finished rows (never null placeholders).
        completed = [r for r in job["results"] if r]
        payload = {
            "status": job["status"],
            "total": job["total"],
            "done": job["done"],
            "counts": dict(job["counts"]),
            "results": completed[-150:],
        }
    return jsonify(payload)


def aligned_rows(headers, data_rows):
    col_count = len(headers)
    for row in data_rows:
        col_count = max(col_count, len(row))
    out_headers = list(headers) + [""] * (col_count - len(headers))
    out_rows = []
    for row in data_rows:
        cells = list(row) + [None] * (col_count - len(row))
        out_rows.append(cells[:col_count])
    return out_headers, out_rows


def send_bytes(data, filename, mimetype):
    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name=filename,
        mimetype=mimetype,
    )


def build_original_file(headers, data_rows, file_ext):
    out_headers, out_rows = aligned_rows(headers, data_rows)
    if file_ext == "csv":
        text = io.StringIO()
        writer = csv.writer(text, lineterminator="\n")
        writer.writerow(out_headers)
        for row in out_rows:
            writer.writerow([csv_cell(c) for c in row])
        return (
            text.getvalue().encode("utf-8-sig"),
            "valid_emails.csv",
            "text/csv",
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Emails"
    ws.append(out_headers)
    for row in out_rows:
        ws.append(row)
    output = io.BytesIO()
    wb.save(output)
    return (
        output.getvalue(),
        "valid_emails.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def export_valid_file(job):
    """Download only VALID rows, keeping the original uploaded file columns."""
    headers = list(job.get("headers") or [])
    source_rows = list(job.get("source_rows") or [])
    ordered = list(job.get("ordered") or [])
    file_ext = job.get("file_ext") or "xlsx"
    valid_count = job.get("counts", {}).get("VALID", 0)

    if not headers:
        return jsonify({"error": "Original columns were not saved. Please upload the file and run verification again."}), 400

    valid_rows = []
    for idx, result in enumerate(ordered):
        if not result or result.get("status") != "VALID":
            continue
        if idx < len(source_rows):
            valid_rows.append(source_rows[idx])
        else:
            valid_rows.append([result.get("email", "")])

    if not valid_rows:
        return jsonify({"error": f"No valid emails to download ({valid_count} valid)."}), 400

    data, filename, mimetype = build_original_file(headers, valid_rows, file_ext)
    return send_bytes(data, filename, mimetype)


def export_all_file(job):
    rows = [r for r in (job.get("ordered") or job["results"]) if r]
    wb = Workbook()
    ws = wb.active
    ws.title = "Verified Emails"
    ws.append(["Email", "Status", "Reason"])
    for item in rows:
        ws.append([item["email"], item["status"], item["reason"]])
    output = io.BytesIO()
    wb.save(output)
    return send_bytes(
        output.getvalue(),
        "verified_emails.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/api/export/<job_id>")
@app.get("/api/export/<job_id>/<kind>")
def export(job_id, kind=None):
    only = (kind or request.args.get("only") or "all").strip().lower()
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            return jsonify({"error": "Job not found. Run verification again."}), 404
        if job["status"] != "completed":
            return jsonify({"error": "Verification is not completed"}), 400
        snapshot = {
            "headers": list(job.get("headers") or []),
            "source_rows": list(job.get("source_rows") or []),
            "ordered": list(job.get("ordered") or []),
            "results": list(job.get("results") or []),
            "file_ext": job.get("file_ext") or "xlsx",
            "counts": dict(job.get("counts") or {}),
        }

    if only == "valid":
        return export_valid_file(snapshot)
    return export_all_file(snapshot)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
